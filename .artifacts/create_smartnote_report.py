from pathlib import Path
from datetime import date
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(r"C:\Users\scara\Downloads\LTTBDD")
OUT = ROOT / ".artifacts" / "Bao_cao_danh_gia_du_an_SmartNote.docx"
IMG = ROOT / ".artifacts" / "report_images"
IMG.mkdir(parents=True, exist_ok=True)

BLUE = "2E74B5"; NAVY = "1F4D78"; LIGHT = "F2F4F7"; RED = "C00000"; GREEN = "2E7D32"

def font(size=24, bold=False):
    candidates = [r"C:\Windows\Fonts\calibrib.ttf" if bold else r"C:\Windows\Fonts\calibri.ttf", r"C:\Windows\Fonts\arial.ttf"]
    for p in candidates:
        if Path(p).exists(): return ImageFont.truetype(p, size)
    return ImageFont.load_default()

def diagram(name, title, boxes, arrows):
    im = Image.new("RGB", (1500, 760), "white"); d = ImageDraw.Draw(im)
    d.text((55, 28), title, fill="#1F4D78", font=font(38, True))
    for key, (x,y,w,h,label,color) in boxes.items():
        d.rounded_rectangle((x,y,x+w,y+h), 18, fill=color, outline="#2E74B5", width=3)
        lines = label.split("\n"); yy=y+(h-len(lines)*30)//2
        for line in lines:
            bb=d.textbbox((0,0),line,font=font(23,True)); d.text((x+(w-(bb[2]-bb[0]))/2,yy),line,fill="#17324D",font=font(23,True)); yy+=31
    for a,b,label in arrows:
        ax,ay,aw,ah,*_=boxes[a]; bx,by,bw,bh,*_=boxes[b]
        p1=(ax+aw/2,ay+ah) if by>ay else (ax+aw,ay+ah/2)
        p2=(bx+bw/2,by) if by>ay else (bx,by+bh/2)
        d.line((p1,p2),fill="#667085",width=4)
        d.polygon([(p2[0],p2[1]),(p2[0]-10,p2[1]-18),(p2[0]+10,p2[1]-18)],fill="#667085")
        if label: d.text(((p1[0]+p2[0])/2+8,(p1[1]+p2[1])/2-24),label,fill="#475467",font=font(18))
    path=IMG/name; im.save(path); return path

arch = diagram("architecture.png","Kiến trúc tổng thể SmartNote",{
 "u":(60,190,260,120,"Người dùng\nAndroid","#E8F1FB"), "app":(440,160,360,170,"Flutter App\nUI · Riverpod · GoRouter","#DCEBFA"),
 "db":(965,100,360,120,"SQLite local\n10 bảng · offline-first","#E9F7EF"), "cloud":(965,300,360,120,"Supabase\nAuth · Postgres · RLS","#F4ECF7"),
 "ext":(965,500,360,120,"Dịch vụ ngoài\nQuote API · Notification","#FFF4E5")},[("u","app","tương tác"),("app","db","đọc/ghi"),("app","cloud","đồng bộ"),("app","ext","HTTP / OS")])

erd = diagram("erd.png","ERD rút gọn — cơ sở dữ liệu cục bộ",{
 "n":(560,80,340,110,"notes\nPK id","#DCEBFA"), "t":(80,260,280,100,"tags\nPK name","#E9F7EF"),
 "nt":(430,260,300,100,"note_tags\nPK note_id, tag_name","#F2F4F7"), "c":(800,260,300,100,"checklist_items\nFK note_id","#F2F4F7"),
 "i":(1140,260,280,100,"note_images\nFK note_id","#F2F4F7"), "r":(110,500,300,100,"note_reminders\nPK/FK note_id","#FFF4E5"),
 "v":(460,500,300,100,"note_versions\nFK note_id","#FFF4E5"), "d":(810,500,300,100,"note_drafts\nPK note_id","#FFF4E5"),
 "o":(1160,500,280,100,"sync_outbox\nnote_id","#F4ECF7")},[("n","nt","1–N"),("t","nt","1–N"),("n","c","1–N"),("n","i","1–N"),("n","r","1–0..1"),("n","v","1–N"),("n","d","1–0..1"),("n","o","1–N")])

seq = diagram("sync_sequence.png","Luồng đồng bộ ghi chú",{
 "u":(70,120,240,100,"Người dùng","#E8F1FB"), "ui":(400,120,260,100,"Flutter UI","#DCEBFA"),
 "repo":(760,120,280,100,"Repository / Outbox","#E9F7EF"), "sql":(1110,120,280,100,"SQLite","#FFF4E5"),
 "sup":(760,500,280,100,"Supabase + RLS","#F4ECF7")},[("u","ui","lưu"),("ui","repo","save"),("repo","sql","transaction"),("repo","sup","push / pull")])

doc=Document(); sec0=doc.sections[0]
sec0.top_margin=Inches(1); sec0.bottom_margin=Inches(1); sec0.left_margin=Inches(1); sec0.right_margin=Inches(1)
sec0.header_distance=Inches(.492); sec0.footer_distance=Inches(.492)
styles=doc.styles
normal=styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(11); normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.10
for nm,size,color,before,after in [('Title',28,NAVY,0,10),('Heading 1',16,BLUE,16,8),('Heading 2',13,BLUE,12,6),('Heading 3',12,NAVY,8,4)]:
    s=styles[nm]; s.font.name='Calibri'; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color); s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after)

def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def set_cell_text(cell,text,bold=False,color=None,size=9):
    cell.text=''; p=cell.paragraphs[0]; r=p.add_run(str(text)); r.bold=bold; r.font.name='Calibri'; r.font.size=Pt(size)
    if color:r.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
def table(headers,rows,widths=None,size=8.5):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,h in enumerate(headers): shade(t.rows[0].cells[i],LIGHT); set_cell_text(t.rows[0].cells[i],h,True,NAVY,size)
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row): set_cell_text(cells[i],v,False,None,size)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(2); return t
def bullet(text): doc.add_paragraph(text,style='List Bullet')
def add_pic(path,caption):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(path),width=Inches(6.45))
    c=doc.add_paragraph(caption); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; c.style='Caption'
def page(): doc.add_page_break()
def field(par,code):
    r=par.add_run(); begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin'); instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text=code; sep=OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'),'separate'); end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end'); r._r.extend([begin,instr,sep,end])

# Header/footer
hp=sec0.header.paragraphs[0]; hp.text='SMARTNOTE  |  BÁO CÁO ĐÁNH GIÁ DỰ ÁN'; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; hp.runs[0].font.size=Pt(8); hp.runs[0].font.color.rgb=RGBColor.from_string('667085')
fp=sec0.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; fp.add_run('SmartNote  •  '); field(fp,'PAGE')

# Cover
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(90)
r=p.add_run('BÁO CÁO ĐÁNH GIÁ\nDỰ ÁN SMARTNOTE'); r.bold=True; r.font.name='Calibri'; r.font.size=Pt(28); r.font.color.rgb=RGBColor.from_string(NAVY)
p=doc.add_paragraph('Phân tích mã nguồn, yêu cầu, kiến trúc, dữ liệu, bảo mật, kiểm thử và triển khai'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].font.size=Pt(14); p.runs[0].font.color.rgb=RGBColor.from_string(BLUE)
doc.add_paragraph('\n')
table(['Thuộc tính','Nội dung'],[['Dự án','SmartNote — ứng dụng ghi chú Flutter offline-first'],['Repository','github.com/theihoz/smartnote'],['Nhánh đánh giá','feature/smartnote-app'],['Phiên bản ứng dụng','1.0.0+1'],['Ngày đánh giá',date.today().strftime('%d/%m/%Y')],['Phạm vi','Mã nguồn cục bộ và cấu hình Supabase production (chỉ đọc)']], [1.6,4.6],10)
doc.add_paragraph('\nTài liệu được lập theo kế hoạch phân tích 28 mục do người dùng cung cấp. Các kết luận phân biệt rõ chức năng đã triển khai, khoảng trống và đề xuất.',style='Quote')
page()

doc.add_heading('Mục lục',0); p=doc.add_paragraph(); field(p,'TOC \\o "1-3" \\h \\z \\u'); doc.add_paragraph('Mở tài liệu trong Microsoft Word và chọn Update Table để cập nhật số trang.',style='Caption'); page()

doc.add_heading('1. Tóm tắt điều hành',0)
doc.add_paragraph('SmartNote là ứng dụng ghi chú Android viết bằng Flutter, ưu tiên hoạt động ngoại tuyến với SQLite và bổ sung đồng bộ tài khoản qua Supabase. Dự án đã có nền tảng chức năng tốt cho một sản phẩm cá nhân: ghi chú văn bản/checklist, tìm kiếm, nhãn, nhắc việc, thùng rác, khóa PIN, xuất PDF/Markdown, song ngữ và giao diện sáng/tối.')
table(['Hạng mục','Đánh giá','Nhận định'],[['Chức năng','Khá','Phần lớn luồng người dùng đã hiện diện; đồng bộ thủ công và lịch sử phiên bản chưa hoàn chỉnh.'],['Kiến trúc','Khá','Tách feature/domain/data/presentation hợp lý; client là modular monolith.'],['Dữ liệu','Khá','SQLite quan hệ tốt; Supabase có RLS, nhưng phạm vi sync hẹp hơn schema.'],['Bảo mật','Trung bình','RLS và quản lý khóa đúng hướng; PIN chỉ bảo vệ UI, chưa mã hóa dữ liệu cục bộ.'],['Kiểm thử','Trung bình–Khá','Có 11 tệp test cục bộ; test và migration hiện không được theo dõi trong Git.'],['Triển khai','Cần cải thiện','Chưa có CI/CD, giám sát và release signing production.']], [1.25,1.25,3.9],9)
doc.add_heading('Kết luận chính',1)
bullet('Phù hợp làm đồ án/sản phẩm MVP và có thể phát triển thành ứng dụng cá nhân đa thiết bị.')
bullet('Rủi ro ưu tiên cao nhất: release Android đang dùng khóa debug; test và migration bị loại khỏi repository; nút đồng bộ chưa thực thi hành động.')
bullet('Supabase production đang ACTIVE_HEALTHY tại khu vực Singapore; 3 bảng public đều bật RLS. Advisor vẫn cảnh báo chính sách có thể áp dụng cho anonymous và tính năng leaked-password protection đang tắt.')

doc.add_heading('2. Tổng quan dự án',0)
doc.add_heading('2.1 Bài toán và mục tiêu',1)
doc.add_paragraph('Dự án giải quyết nhu cầu ghi chép nhanh, tổ chức công việc và truy cập dữ liệu khi không có mạng. Mục tiêu kỹ thuật là phản hồi nhanh nhờ dữ liệu cục bộ, sau đó đồng bộ có kiểm soát lên cloud khi người dùng đăng nhập.')
table(['Nội dung','Kết quả phân tích'],[['Người dùng mục tiêu','Cá nhân, sinh viên, nhân viên văn phòng cần ghi chú/nhắc việc trên Android.'],['Phạm vi','Ứng dụng Android; không có web admin, backend tùy chỉnh hoặc iOS trong repository.'],['Trạng thái','Đang phát triển tích cực; nhiều commit tính năng/sửa lỗi gần đây.'],['License','Không tìm thấy tệp LICENSE — quyền sử dụng/phân phối chưa được tuyên bố rõ.'],['Quản trị','Không có vai trò Administrator trong luồng nghiệp vụ hiện tại.']], [1.5,4.7],9)
doc.add_heading('2.2 Tính năng hiện có',1)
for x in ['Ghi chú văn bản và checklist; ảnh từ camera/thư viện; yêu thích, nhãn, tìm kiếm và lọc.','Tự động lưu bản nháp; lưu tối đa 20 snapshot phiên bản cho mỗi ghi chú.','Nhắc việc bằng thông báo cục bộ; thùng rác soft-delete 30 ngày.','Khóa ghi chú bằng PIN chung 4–6 số; xuất/chia sẻ PDF hoặc Markdown.','Đăng ký/đăng nhập email–mật khẩu; đồng bộ ghi chú qua Supabase.','Tiếng Việt/Anh và light/dark/system được lưu bằng SharedPreferences.']: bullet(x)

doc.add_heading('3. Repository và luồng chạy',0)
table(['Khu vực','Vai trò'],[['lib/main.dart','Entry point; khởi tạo SharedPreferences, SQLite, notification, Supabase tùy cấu hình và Riverpod.'],['lib/app','Theme, routing, provider và shell điều hướng responsive.'],['lib/features','Các feature auth, notes, reminders, security, export, sync, settings, inspiration.'],['lib/l10n','ARB/generated localization và helper featureText.'],['android','Manifest, Gradle và cấu hình build Android.'],['test','11 tệp test cục bộ; hiện bị .gitignore loại khỏi repository.'],['supabase/migrations','Migration cloud cục bộ; hiện bị .gitignore loại khỏi repository.'],['download_site','Trang tải APK cục bộ; hiện không theo dõi Git.']], [1.55,4.65],8.5)
doc.add_heading('Luồng khởi động',1)
doc.add_paragraph('Flutter binding → tải cài đặt → mở SQLite v5 → khởi tạo repository/scheduler → thử cấu hình Supabase từ --dart-define → nếu có session thì xử lý outbox và kéo dữ liệu cloud → dựng MaterialApp.router. Nếu cloud lỗi, ứng dụng vẫn tiếp tục ở chế độ local; tuy nhiên lỗi bị nuốt và không có telemetry.')
add_pic(arch,'Hình 1 — Kiến trúc tổng thể và các hệ thống liên quan')

doc.add_heading('4. Technology stack',0)
rows=[('Client','Flutter / Dart','SDK Dart ^3.12.2','UI đa nền tảng','Material 3; repository hiện chỉ có Android'),('State','flutter_riverpod','2.6.1','Quản lý trạng thái','Provider override hỗ trợ kiểm thử'),('Routing','go_router','17.3.0','Điều hướng','Route rõ theo feature'),('Local DB','SQLite / sqflite','2.4.3','Offline-first','Transaction + FK; có một index tạo mới thiếu'),('Cloud','Supabase Flutter','2.16.0','Auth/Postgres sync','Dùng publishable key qua dart-define'),('HTTP','http','1.6.0','Quote API','Không có backend REST riêng'),('Export','pdf / share_plus','3.13.0 / 13.3.0','PDF/Markdown','Xử lý trên thiết bị'),('Notification','flutter_local_notifications','22.2.0','Nhắc việc','Có timezone'),('Settings','shared_preferences','2.5.5','Theme/locale/PIN','Không phù hợp lưu bí mật mạnh'),('CI/CD','Không có','—','—','Khoảng trống triển khai')]
table(['Thành phần','Công nghệ','Phiên bản','Vai trò','Nhận xét'],rows,[1.0,1.35,1.0,1.25,2.0],7.5)

doc.add_heading('5. Actor và yêu cầu',0)
table(['Actor','Quyền hạn','Dữ liệu truy cập'],[['Guest / Local user','Dùng toàn bộ chức năng cục bộ, không đồng bộ cloud.','SQLite và file ảnh trên thiết bị.'],['Authenticated user','Như local user; đăng nhập và đồng bộ ghi chú của chính mình.','Bản ghi có user_id = auth.uid() trên Supabase.'],['Android OS','Cấp quyền camera/thông báo; kích hoạt notification.','Quyền thiết bị trong phạm vi ứng dụng.'],['Supabase','Xác thực và lưu dữ liệu cloud theo RLS.','notes, note_reminders, note_versions.']], [1.25,2.8,2.4],8.5)
fr=[('FR-001','Tạo/sửa ghi chú','User','Lưu text hoặc checklist, nhãn, ảnh, màu','High'),('FR-002','Tìm kiếm/lọc','User','Tìm theo nội dung, yêu thích, nhãn và sắp xếp','High'),('FR-003','Hoàn tác','User','Hoàn tác tạo/sửa/xóa trong 5 giây','High'),('FR-004','Bản nháp','User','Tự động lưu nội dung đang soạn','High'),('FR-005','Phiên bản','User','Lưu tối đa 20 snapshot mỗi ghi chú','Medium'),('FR-006','Nhắc việc','User','Lập lịch và thông báo cục bộ','High'),('FR-007','Thùng rác','User','Soft-delete, khôi phục trong 30 ngày','High'),('FR-008','Khóa ghi chú','User','Bảo vệ truy cập UI bằng PIN','High'),('FR-009','Xuất/chia sẻ','User','Gộp nhiều ghi chú thành PDF/Markdown','Medium'),('FR-010','Tài khoản','User','Đăng ký/đăng nhập email–mật khẩu','High'),('FR-011','Đồng bộ','Auth user','Push outbox, pull bản cập nhật theo thời gian','High'),('FR-012','Giao diện','User','VI/EN và light/dark/system có lưu','Medium')]
doc.add_heading('5.1 Functional Requirements',1); table(['ID','Chức năng','Actor','Mô tả','Ưu tiên'],fr,[.7,1.15,1.0,3.0,.7],7.7)
nfr=[('NFR-001','Performance','Mở/list dữ liệu cục bộ phản hồi nhanh; cần đo p95.'),('NFR-002','Reliability','Mất cloud không làm mất khả năng ghi chú local.'),('NFR-003','Security','Mọi bảng exposed bật RLS; không nhúng service-role key.'),('NFR-004','Privacy','Dữ liệu local nhạy cảm cần mã hóa at-rest trong bản production.'),('NFR-005','Availability','Cloud phụ thuộc Supabase; local vẫn hoạt động offline.'),('NFR-006','Maintainability','Feature modules, lint sạch, test/migration phải theo dõi Git.'),('NFR-007','Usability','Luồng chính tối đa ít thao tác, phản hồi lỗi rõ ràng.'),('NFR-008','Accessibility','Semantics, tương phản, kích thước chạm và text scaling phải được kiểm tra.'),('NFR-009','Compatibility','Android minSdk 24; kiểm thử nhiều kích cỡ và phiên bản OS.'),('NFR-010','Scalability','Sync phân trang/batch, retry backoff khi dữ liệu tăng.')]
doc.add_heading('5.2 Non-functional Requirements',1); table(['ID','Nhóm','Tiêu chí'],nfr,[.8,1.3,4.5],8)

page(); doc.add_heading('6. Phân tích nghiệp vụ và UML rút gọn',0)
doc.add_heading('6.1 Danh sách use case',1)
doc.add_paragraph('UC-01 Quản lý ghi chú; UC-02 Tìm kiếm/lọc; UC-03 Lập nhắc việc; UC-04 Xóa/khôi phục; UC-05 Khóa/mở khóa; UC-06 Xuất/chia sẻ; UC-07 Đăng ký/đăng nhập; UC-08 Đồng bộ cloud; UC-09 Đổi ngôn ngữ/giao diện.')
doc.add_heading('6.2 Đặc tả UC-01 — Tạo hoặc chỉnh sửa ghi chú',1)
table(['Thuộc tính','Nội dung'],[['Actor','Local user / Authenticated user'],['Tiền điều kiện','Ứng dụng đã mở; SQLite khả dụng.'],['Trigger','Người dùng chọn tạo mới hoặc mở ghi chú.'],['Luồng chính','Nhập dữ liệu → validator kiểm tra → repository lưu transaction → tạo snapshot/outbox → tải lại danh sách → hiển thị Snackbar hoàn tác.'],['Luồng thay thế','Nội dung đang soạn được lưu draft; có thể thêm checklist, ảnh, nhãn, reminder.'],['Ngoại lệ','Validation thất bại hoặc SQLite lỗi: giữ màn hình và hiển thị lỗi.'],['Hậu điều kiện','Ghi chú local nhất quán; sẵn sàng sync khi có tài khoản.']], [1.35,4.85],8.5)
doc.add_heading('6.3 Đặc tả UC-08 — Đồng bộ cloud',1)
table(['Thuộc tính','Nội dung'],[['Actor','Authenticated user'],['Tiền điều kiện','Supabase được cấu hình; có session.'],['Trigger hiện tại','Khởi động ứng dụng khi đã có session.'],['Luồng chính','Đọc outbox → upsert/soft-delete cloud → kéo notes mới hơn last_synced_at → áp dụng nếu remote.updated_at mới hơn local.'],['Ngoại lệ','Tăng attempts; giữ outbox. Không có backoff/max retry/dead-letter.'],['Khoảng trống','Đăng nhập không kích hoạt sync ngay; nút Sync trong Settings chưa có callback thực. Reminder/version chưa được sync.']], [1.35,4.85],8.5)
add_pic(seq,'Hình 2 — Sequence rút gọn của luồng lưu và đồng bộ')
doc.add_heading('6.4 Class/component model',1)
doc.add_paragraph('Presentation screens phụ thuộc Riverpod providers; NotesController điều phối NoteRepository; SqliteNoteRepository chịu trách nhiệm transaction/hydration/outbox; OutboxSyncService phối hợp local repository với CloudNoteStore; SupabaseCloudNoteStore gọi Data API. Các service độc lập xử lý draft, reminder, PIN và export. Đây là phân tách trách nhiệm hợp lý, chưa cần tách thành nhiều package.')

doc.add_heading('7. Cơ sở dữ liệu',0)
add_pic(erd,'Hình 3 — ERD rút gọn của SQLite (sync_state không thể hiện)')
doc.add_heading('7.1 Data dictionary rút gọn',1)
dbrows=[('notes','id; title; body; kind; is_favorite; color_key; deleted_at; is_locked; timestamps','PK id','Thực thể chính'),('tags','name','PK, NOCASE','Danh mục nhãn'),('note_tags','note_id; tag_name; position','PK kép; 2 FK','Quan hệ N–N'),('checklist_items','id; note_id; text; is_done; position','PK; FK cascade','Dòng checklist'),('note_images','note_id; path; position','PK kép; FK cascade','Đường dẫn ảnh local'),('note_reminders','note_id; schedule; repeat…','PK/FK cascade','Cấu hình nhắc'),('note_versions','id; note_id; snapshot; created_at','PK; FK cascade; index','Snapshot JSON'),('note_drafts','note_id; title; body; updated_at','PK','Bản nháp'),('sync_outbox','id; note_id; operation; attempts','PK auto','Hàng đợi đồng bộ'),('sync_state','key; last_synced_at','PK key','Mốc pull cloud')]
table(['Bảng','Cột chính','Ràng buộc','Vai trò'],dbrows,[1.25,2.5,1.4,1.4],7.3)
doc.add_heading('7.2 Đánh giá thiết kế dữ liệu',1)
bullet('Ưu điểm: FK được bật; child rows cascade; transaction khi lưu; version giới hạn 20; soft-delete hỗ trợ khôi phục.')
bullet('Lỗi schema: notes_deleted_at_idx chỉ được tạo trong đường upgrade từ v1, không được tạo khi cài mới trực tiếp schema v5.')
bullet('Hiệu năng: hydrate mỗi note truy vấn riêng tags/checklist/images, tạo mẫu N+1; nên batch theo danh sách id khi dữ liệu tăng.')
bullet('Cloud dùng JSONB cho tags/images/checklist để đơn giản sync. Hợp lý cho MVP nhưng giảm khả năng truy vấn quan hệ trên server.')
bullet('Cloud production có 3 bảng: notes, note_reminders, note_versions; tất cả RLS enabled và hiện chưa có dữ liệu.')

doc.add_heading('8. Kiến trúc hệ thống và C4',0)
table(['Mức','Mô tả'],[['C4 L1 — Context','Người dùng tương tác SmartNote; ứng dụng dùng Android OS, Supabase và Quote API.'],['C4 L2 — Container','Flutter Android app + SQLite local + Supabase Auth/Postgres; không có backend riêng.'],['C4 L3 — Component','Presentation → Riverpod/controller → repository/service → SQLite/Supabase/plugin.'],['Deployment','APK trên thiết bị; DB local trong sandbox app; cloud tại ap-southeast-1.']], [1.5,4.7],9)
doc.add_heading('Data flow',1)
doc.add_paragraph('Dữ liệu người dùng được ghi vào SQLite trước. Repository cập nhật các bảng con và outbox trong transaction. Đồng bộ đẩy dữ liệu sang Supabase bằng JWT session; RLS lọc theo user_id. Ảnh hiện chỉ lưu đường dẫn local, vì vậy không thể xuất hiện trên thiết bị thứ hai nếu không bổ sung Storage/upload mapping.')

doc.add_heading('9. API và dịch vụ ngoài',0)
doc.add_paragraph('SmartNote không định nghĩa REST backend tùy chỉnh. “API” của hệ thống là Supabase Auth/Data API do SDK Dart gọi và một Quote API công khai. Vì vậy không có các endpoint /api/... do dự án tự sở hữu để tài liệu hóa.')
table(['Tác vụ','Giao thức/đích','Auth','Request chính','Kết quả/lỗi'],[['Đăng ký','Supabase Auth signUp','Publishable key','email, password','User/session hoặc AuthException'],['Đăng nhập','Supabase Auth signInWithPassword','Publishable key','email, password','Session/JWT hoặc AuthException'],['Đẩy note','Data API: upsert notes','Bearer JWT','Bản ghi note + user_id','Row upsert; RLS bảo vệ'],['Xóa mềm','Data API: update notes','Bearer JWT','deleted_at, updated_at','Row update'],['Kéo note','Data API: select notes','Bearer JWT','updated_at > checkpoint','Danh sách JSON'],['Quote','HTTP GET dịch vụ ngoài','Không','Không/locale tùy repository','Quote hoặc fallback/error']], [1.0,1.55,1.0,1.55,1.65],7.3)
doc.add_heading('Quy ước lỗi cần chuẩn hóa',1)
doc.add_paragraph('UI hiện thường hiển thị chuỗi exception trực tiếp. Nên ánh xạ lỗi thành mã nội bộ (AUTH_INVALID_CREDENTIALS, NETWORK_UNAVAILABLE, SYNC_CONFLICT, STORAGE_FAILURE), thông điệp VI/EN thân thiện và log kỹ thuật riêng; không hiển thị chi tiết nhạy cảm từ backend.')

doc.add_heading('10. Phân tích bảo mật',0)
security=[('High','Release Android dùng debug signing key','APK production không có danh tính phát hành an toàn.','Thiết lập keystore release qua biến bí mật; chặn release nếu thiếu.'),('High','test và Supabase migrations bị ignore','Không tái lập được kiểm thử/schema; tăng rủi ro drift và mất lịch sử.','Theo dõi lại hai thư mục, chỉ ignore artefact sinh ra/secrets.'),('Medium','Advisor: policy có thể áp dụng anonymous','Anonymous sign-in có thể làm thay đổi mô hình tin cậy dù predicate sở hữu vẫn tồn tại.','Tắt anonymous nếu không dùng; policy TO authenticated + auth.uid ownership.'),('Medium','Leaked password protection tắt','Cho phép mật khẩu đã lộ.','Bật trong Supabase Auth và quy định độ mạnh mật khẩu.'),('Medium','PIN hash trong SharedPreferences','Khóa UI, không mã hóa nội dung/DB; thiết bị bị compromise có thể đọc dữ liệu.','Dùng Android Keystore/biometric và SQLCipher nếu dữ liệu nhạy cảm.'),('Medium','Ảnh chỉ là đường dẫn local','Có thể rò dữ liệu qua backup/quyền file tùy vị trí.','Lưu app-private; kiểm MIME/kích thước; dùng Storage RLS khi sync.'),('Low','Exception cloud bị nuốt/hiển thị trực tiếp','Thiếu audit; có thể lộ chi tiết kỹ thuật.','Structured logging + thông điệp lỗi chuẩn hóa.'),('Info','SQL injection/XSS/CSRF','sqflite dùng tham số; app native không có DOM/cookie web.','Duy trì validation; đánh giá lại nếu thêm WebView/web client.')]
table(['Mức','Phát hiện','Ảnh hưởng','Khuyến nghị'],security,[.65,1.55,2.0,2.25],7.1)
doc.add_heading('Cơ chế tốt đang có',1)
bullet('Không nhúng service-role key; cấu hình cloud qua dart-define và publishable key.')
bullet('RLS bật trên cả 3 bảng public; FK user_id tham chiếu auth.users.')
bullet('PIN dùng salt ngẫu nhiên, HMAC-SHA256 120.000 vòng, so sánh constant-time và khóa 30 giây sau 5 lần sai.')
bullet('Input email/password/PIN/note được kiểm tra; thao tác SQLite dùng API tham số hóa.')

doc.add_heading('11. Kiểm thử',0)
doc.add_paragraph('Repository cục bộ có 11 tệp test bao phủ domain, NotesController, SQLite repositories, draft, reminder, PIN, export, quote, sync và widget app. Lần chạy xác nhận trước trong quá trình phát triển ghi nhận 48 test pass, flutter analyze không có lỗi và debug APK build thành công. Trong lần lập báo cáo này, lệnh kiểm thử không kết thúc và không phát sinh output trong thời gian chờ, nên không coi đó là bằng chứng chạy mới; cần kiểm tra môi trường Flutter/process lock.')
table(['Loại','Hiện trạng','Khoảng trống'],[['Unit','Có','Cần giữ trong Git và CI.'],['Widget/UI','Có smartnote_app_test','Chưa có regression ảnh/accessibility.'],['Integration','Chưa thấy test thật trong integration_test','Thêm luồng DB + notification + auth mock.'],['API/live','Không','Không nên phụ thuộc production; dùng project test.'],['Security','Một số PIN/RLS gián tiếp','Thiếu kiểm tra policy và secret scanning.'],['Performance','Không','Đo cold start, list 1k notes, sync batch.'],['UAT','Không','Lập checklist thiết bị thật VI/EN, light/dark.']], [1.15,2.25,2.9],8)
tc=[('TC-001','Tạo note hợp lệ','Title/body hợp lệ','Lưu, xuất hiện danh sách','Đã có test','Pass trước'),('TC-002','Undo xóa','Xóa rồi undo <5s','Note khôi phục','Đã có test','Pass trước'),('TC-003','Draft','Thoát editor sau nhập','Mở lại còn draft','Đã có test','Pass trước'),('TC-004','PIN lockout','Sai PIN 5 lần','Khóa 30 giây','Đã có test','Pass trước'),('TC-005','Sync conflict','Remote mới hơn local','Remote thắng','Đã có test','Pass trước'),('TC-006','Theme persistence','Chọn dark, khởi động lại','Vẫn dark','Có settings test','Pass trước'),('TC-007','Reminder OS','Lịch tương lai','Notification đúng giờ','Chưa chạy thiết bị','Pending'),('TC-008','RLS isolation','User A đọc note B','0 row/denied','Chưa tự động hóa','Pending')]
doc.add_heading('Test cases đại diện',1); table(['ID','Feature','Input','Expected','Actual','Status'],tc,[.65,1.1,1.4,1.55,1.25,.75],7.2)

doc.add_heading('12. Deployment',0)
table(['Hạng mục','Hiện trạng','Đánh giá'],[['Development','Flutter SDK; pubspec/lockfile; Android minSdk 24','Có thể dựng local khi cấu hình SDK.'],['Environment','SUPABASE_URL và SUPABASE_PUBLISHABLE_KEY qua --dart-define','Đúng hướng; cần tài liệu hóa template không chứa secret.'],['Build','flutter build apk','Release hiện ký bằng debug key — không dùng production.'],['Database migration','3 migration đã áp dụng trên Supabase production','Tệp migration local bị ignore, không tái lập được.'],['Hosting','Supabase ap-southeast-1; trang tải/APK từng được mô tả','Không có pipeline phát hành tự động.'],['CI/CD','Không có','Cần GitHub Actions: format/analyze/test/build/secret scan.'],['Backup/monitoring','Không cấu hình trong repo','Cần PITR/backup theo gói, crash reporting và sync metrics.'],['HTTPS','Supabase/Quote API qua HTTPS','Được cung cấp bởi dịch vụ; không có reverse proxy.']], [1.3,3.0,2.05],8.2)
doc.add_heading('Deployment guide tối thiểu',1)
for x in ['Cài Flutter tương thích và Android SDK; chạy flutter pub get.','Cấp SUPABASE_URL và publishable key bằng --dart-define; tuyệt đối không dùng service-role key.','Chạy format, analyze và test; kiểm tra migration production khớp repository.','Tạo keystore release, lưu secrets ngoài Git; build appbundle/APK release.','Kiểm thử trên thiết bị Android thật: permission, notification, offline/sync, VI/EN và light/dark.','Phát hành qua kênh kiểm thử; theo dõi crash, auth và sync trước production.']: bullet(x)

doc.add_heading('13. Chất lượng mã và kiến trúc',0)
table(['Tiêu chí','Nhận xét','Mức'],[['Cấu trúc','Feature-based, domain/data/presentation rõ','Tốt'],['Naming','Tên lớp/tệp nhất quán, dễ truy vết','Tốt'],['Separation of concerns','Repository/service tách khỏi UI; controller tương đối gọn','Tốt'],['Error handling','Có fallback offline nhưng nuốt lỗi cloud và UI đôi khi lộ exception','Cần cải thiện'],['Logging/monitoring','Hầu như chưa có','Yếu'],['Testing','Phạm vi unit/widget tốt cho MVP nhưng không được commit','Trung bình'],['Documentation','README có hướng dẫn cơ bản; thiếu license/API/deployment chuẩn','Trung bình'],['Dependencies','Phiên bản pin trong pubspec.lock; bộ thư viện phù hợp','Khá'],['Performance','Có N+1 hydration và sync chưa batch/backoff','Cần đo']], [1.55,3.9,.9],8.2)
doc.add_heading('Điểm mạnh',1)
bullet('Offline-first thực tế, transaction và outbox giảm nguy cơ mất dữ liệu khi mạng gián đoạn.')
bullet('UI/UX đã có nền tảng Material 3, responsive, song ngữ và lưu theme.')
bullet('RLS theo người dùng, xác thực chuẩn Supabase và không đưa khóa đặc quyền vào client.')
doc.add_heading('Điểm yếu',1)
bullet('Cloud schema và client sync chưa đồng nhất: reminder/version tồn tại nhưng sync service chỉ xử lý notes.')
bullet('Sync chủ yếu lúc khởi động; nút Sync chưa hoạt động và đăng nhập thành công chưa kéo dữ liệu ngay.')
bullet('History snapshot chưa có màn hình xem/khôi phục; ảnh không đồng bộ đa thiết bị.')

doc.add_heading('14. Technical debt và Improvement plan',0)
debt=[('TD-01','Release debug-signed','High','Không phát hành an toàn','Release keystore + CI guard'),('TD-02','Test/migration bị ignore','High','Mất khả năng tái lập','Đưa lại vào Git'),('TD-03','Nút Sync rỗng','High','Người dùng hiểu sai trạng thái','Gọi sync service + trạng thái tiến trình'),('TD-04','Reminder/version không sync','Medium','Dữ liệu khác nhau giữa thiết bị','Mở rộng outbox hoặc ghi rõ local-only'),('TD-05','N+1 SQLite hydration','Medium','Chậm khi nhiều notes','Batch query theo note_id'),('TD-06','LWW dựa timestamp','Medium','Xung đột/clock skew','Server timestamp + conflict metadata'),('TD-07','PIN/DB chưa mã hóa','Medium','Rò dữ liệu khi thiết bị compromise','Keystore/biometric + encrypted DB'),('TD-08','Thiếu observability','Medium','Khó chẩn đoán lỗi production','Crash/log/sync metrics'),('TD-09','Hai cơ chế localization','Low','Dễ lệch bản dịch','Quy về ARB/generated l10n'),('TD-10','Index deleted_at thiếu khi cài mới','Low','Query trash có thể chậm','Tạo index trong _createSchema')]
table(['ID','Vấn đề','Mức','Ảnh hưởng','Đề xuất'],debt,[.65,1.7,.7,1.65,2.0],7.2)
doc.add_heading('Short-term — 1 đến 2 tuần',1)
for x in ['Sửa release signing; đưa test và migrations trở lại Git; thêm CI tối thiểu.','Nối nút Sync và sync sau đăng nhập; hiển thị trạng thái/thời điểm sync gần nhất.','Sửa index fresh-install; chuẩn hóa lỗi; bật leaked-password protection và rà policy anonymous.']: bullet(x)
doc.add_heading('Medium-term — 1 đến 2 tháng',1)
for x in ['Đồng bộ reminder/version hoặc tuyên bố rõ local-only; bổ sung Storage cho ảnh.','Tối ưu batch hydration và sync retry/backoff; thêm integration/UAT/accessibility tests.','Màn hình lịch sử phiên bản và khôi phục; telemetry không chứa nội dung ghi chú.']: bullet(x)
doc.add_heading('Long-term — 3 đến 6 tháng',1)
for x in ['Mã hóa dữ liệu local, biometric/Keystore và threat model đầy đủ.','Đồng bộ xung đột rõ ràng hơn, phân trang/batch và hỗ trợ nhiều thiết bị.','Quy trình phát hành production, backup/khôi phục và SLA/monitoring.']: bullet(x)

doc.add_heading('15. Roadmap đề xuất',0)
table(['Giai đoạn','Mục tiêu','Điều kiện hoàn thành'],[['Phase 1 — Stabilize','Signing, Git hygiene, CI, Sync UX, security advisor','Build release có chữ ký; CI xanh; migration/test tái lập được.'],['Phase 2 — Complete','Version UI, cloud parity, ảnh, integration test','Dữ liệu nhất quán hai thiết bị; test luồng chính tự động.'],['Phase 3 — Harden & Scale','Encryption, conflict model, monitoring, performance','Đạt ngưỡng bảo mật/hiệu năng đã định lượng.']], [1.35,2.45,2.55],8.2)

doc.add_heading('16. Checklist và trả lời cuối cùng',0)
done=['Hiểu mục tiêu, phạm vi và người dùng','Phân tích repository, stack và luồng chạy','Xác định actor, FR và NFR','Lập use case/sequence/class model rút gọn','Phân tích SQLite, Supabase, ERD và data dictionary','Phân tích architecture/C4/data flow','Phân tích API, security, testing và deployment','Đánh giá code/architecture, technical debt và roadmap']
table(['Trạng thái','Hạng mục'],[['Hoàn thành',x] for x in done],[1.1,5.2],8.5)
qa=[('1. Dự án giải quyết gì?','Ghi chú và nhắc việc cá nhân, ưu tiên offline và có đồng bộ.'),('2. Ai sử dụng?','Người dùng local và người dùng đã xác thực; không có admin app.'),('3. Chức năng?','CRUD note/checklist, search/tag, draft/version, reminder, trash, lock, export, sync, VI/EN/theme.'),('4. Source tổ chức?','Feature-based Flutter, phân lớp presentation/application/domain/data.'),('5. Công nghệ?','Flutter/Dart, Riverpod, GoRouter, SQLite, Supabase, notification/PDF plugins.'),('6. Database?','10 bảng SQLite quan hệ; 3 bảng Supabase JSONB + RLS.'),('7. Thành phần giao tiếp?','UI → controller/provider → repository/service → SQLite/Supabase/OS.'),('8. API?','Supabase Auth/Data API và Quote API; không có REST backend riêng.'),('9. Bảo mật?','RLS, JWT, publishable key, PIN hashing; còn thiếu encryption và hardening production.'),('10. Kiểm thử?','11 tệp unit/widget local; thiếu CI, integration, security/performance/UAT.'),('11. Triển khai?','APK Android + Supabase Singapore; thiếu release signing và CI/CD.'),('12. Điểm mạnh/yếu?','Mạnh offline-first/modular; yếu ở cloud parity, observability và release process.'),('13. Cần cải thiện?','Git hygiene, sync UX, signing, security advisor, test automation, performance.'),('14. Hướng phát triển?','Hoàn thiện đa thiết bị, ảnh cloud, history UI, encryption và monitoring.')]
table(['Câu hỏi','Trả lời'],qa,[2.15,4.15],7.7)

doc.add_heading('17. Phạm vi và nguồn bằng chứng',0)
doc.add_paragraph('Đánh giá dựa trên mã nguồn tại workspace, README/pubspec, lịch sử Git, schema SQLite, test cục bộ và truy vấn chỉ đọc Supabase project ijfiountpfumysfbwrqh ngày lập báo cáo. Không thực hiện penetration test, kiểm thử tải, kiểm thử thiết bị thật hoặc thay đổi dữ liệu/schema production. Issue/PR trên GitHub không được dùng làm bằng chứng vì không có dữ liệu cục bộ đáng tin cậy trong phạm vi đánh giá.')
doc.add_paragraph('Mức độ ưu tiên là đánh giá kỹ thuật tại thời điểm hiện tại, không thay thế kiểm toán bảo mật độc lập hoặc tiêu chuẩn phát hành của cửa hàng ứng dụng.',style='Quote')

doc.core_properties.title='Báo cáo đánh giá dự án SmartNote'
doc.core_properties.subject='Phân tích phần mềm theo kế hoạch 28 mục'
doc.core_properties.author='Codex'
doc.core_properties.keywords='SmartNote, Flutter, Supabase, SQLite, SRS, UML, Security, Testing'
OUT.parent.mkdir(parents=True,exist_ok=True); doc.save(OUT)
print(OUT)
